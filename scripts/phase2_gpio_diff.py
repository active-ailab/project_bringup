#!/usr/bin/env python3
"""Phase2: dynamically locate GPIO table structure and emit BASE≠NEW diffs.

Does NOT assume a fixed form A/B/C. Those are only historical examples.
Agent / script must discover: header row, pin-id columns, product columns.

Examples:
  python3 scripts/phase2_gpio_diff.py \\
    --xlsx gpio.xlsx --base-col OSLO --new-col TestOslo

  python3 scripts/phase2_gpio_diff.py \\
    --xlsx gpio.xlsx --sheet Ap330BP-GPIO \\
    --base-col Meru --new-col TestMeru

  python3 scripts/phase2_gpio_diff.py \\
    --json-sheet export.json --base-col Rimo --new-col testRimo
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple


# Column headers that often identify a pin / pad / net (not exhaustive).
PIN_HINTS = (
    "pin name",
    "pinname",
    "gpio name",
    "gpio",
    "ball",
    "net name",
    "net",
    "signal",
    "io name",
    "脚名",
    "管脚",
    "引脚",
)
PAD_HINTS = ("pad", "pad name", "pad no", "pad#", "焊盘")
PKG_PIN_HINTS = ("pin", "pkg pin", "package pin", "ball", "封装脚")


@dataclass
class Schema:
    """Discovered table schema (dynamic; not a fixed A/B/C enum)."""

    header_row: int
    headers: List[str]
    base_col: int
    new_col: int
    id_cols: List[Tuple[str, int]]  # (role, index) e.g. ("pin", 0), ("pad", 1)
    sheet: Optional[str] = None
    notes: List[str] = field(default_factory=list)

    def kind_hint(self) -> str:
        roles = {r for r, _ in self.id_cols}
        if "pad" in roles and ("pin" in roles or "pkg_pin" in roles):
            return "pin+pad"
        if "pin" in roles:
            return "pin-name"
        return "custom"


@dataclass
class DiffRow:
    ids: Dict[str, str]
    base_sig: str
    new_sig: str
    kind: str = "functional"  # functional | alias

    def as_dict(self) -> dict:
        return {
            "ids": self.ids,
            "base": self.base_sig,
            "new": self.new_sig,
            "kind": self.kind,
        }

    def label(self) -> str:
        parts = [f"{k}={v}" for k, v in self.ids.items() if v]
        return "/".join(parts) if parts else "?"


class SchemaError(Exception):
    """Recoverable schema discovery failure (e.g. wrong sheet while probing)."""


def die(msg: str) -> None:
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)

def norm(s: Any) -> str:
    if s is None:
        return ""
    return str(s).replace("\n", " ").strip()


def norm_header(s: Any) -> str:
    return re.sub(r"\s+", " ", norm(s)).lower()


def col_match(headers: Sequence[str], want: str) -> Optional[int]:
    want_l = want.strip().lower()
    if not want_l:
        return None
    want_compact = re.sub(r"[^a-z0-9]", "", want_l)
    for i, h in enumerate(headers):
        if h and h == want_l:
            return i
    for i, h in enumerate(headers):
        if h and (h.startswith(want_l) or want_l in h):
            return i
    # tolerate product name vs column display (rimo vs Rimo / testRimo)
    # never match empty/tiny headers: "" in "rimo2rocky" is always true
    if len(want_compact) < 2:
        return None
    for i, h in enumerate(headers):
        hc = re.sub(r"[^a-z0-9]", "", h or "")
        if len(hc) < 2:
            continue
        if hc == want_compact or want_compact in hc:
            return i
    return None


def score_header_row(row: Sequence[Any], base_col: str, new_col: str) -> int:
    """Higher = more likely a GPIO product matrix header."""
    cells = [norm_header(c) for c in row if norm(c)]
    if len(cells) < 3:
        return 0
    score = 0
    joined = " ".join(cells)
    for hint in PIN_HINTS + PAD_HINTS + PKG_PIN_HINTS:
        if any(hint == c or hint in c for c in cells):
            score += 3
    if col_match(cells, base_col) is not None:
        score += 8
    if col_match(cells, new_col) is not None:
        score += 8
    # many short product-like headers
    if sum(1 for c in cells if 2 <= len(c) <= 24) >= 4:
        score += 2
    # penalize pure description rows
    if "说明" in joined or "note" in joined and score < 8:
        score -= 2
    return score


def find_best_header(
    rows: List[Sequence[Any]], base_col: str, new_col: str
) -> Tuple[int, List[str]]:
    best_i, best_s = -1, -1
    for i, r in enumerate(rows[:20]):
        s = score_header_row(r, base_col, new_col)
        if s > best_s:
            best_s, best_i = s, i
    if best_i < 0 or best_s < 8:
        preview = []
        for i, r in enumerate(rows[:8]):
            preview.append(f"  row{i}: {[norm(c) for c in r[:12] if norm(c)]}")
        raise SchemaError(
            "cannot locate GPIO header with BASE/NEW product columns; "
            f"score={best_s}. preview:\n" + "\n".join(preview)
        )
    headers = [norm_header(c) for c in rows[best_i]]
    return best_i, headers


def _is_meta_header(h: str) -> bool:
    return any(
        k in h
        for k in (
            "function",
            "fuction",
            "type",
            "描述",
            "说明",
            "remark",
            "comment",
            "接口功能",
            "复用",
            "default",
            "digital",
        )
    )


def classify_id_columns(headers: List[str], base_i: int, new_i: int) -> List[Tuple[str, int]]:
    """Pick columns that identify a row (pin/pad/net…), excluding product cols."""
    skip = {base_i, new_i}
    found: List[Tuple[str, int]] = []

    def take(role: str, idx: int) -> None:
        if idx in skip or any(i == idx for _, i in found):
            return
        found.append((role, idx))

    for i, h in enumerate(headers):
        if i in skip or not h or _is_meta_header(h):
            continue
        if any(h == x or h.startswith(x + " ") for x in PAD_HINTS) or h == "pad":
            take("pad", i)
        elif "pin name" in h or h in ("pinname", "gpio name", "io name", "脚名", "管脚", "引脚"):
            take("pin", i)
        elif h == "gpio" or (h.startswith("gpio") and "type" not in h):
            take("gpio", i)
        elif h in ("net", "net name", "signal") or "net name" in h:
            take("net", i)

    has_pad = any(r == "pad" for r, _ in found)
    has_pin = any(r == "pin" for r, _ in found)
    if not has_pin:
        for i, h in enumerate(headers):
            if i in skip or _is_meta_header(h):
                continue
            if h == "pin" or h in PKG_PIN_HINTS:
                take("pkg_pin" if has_pad else "pin", i)
                break

    if not found:
        for i, h in enumerate(headers):
            if i in skip or not h or _is_meta_header(h):
                continue
            take("id", i)
            break

    if not found:
        raise SchemaError(f"no pin/pad/id column found in headers={headers[:24]}")
    return found

def load_xlsx_all_sheets(path: Path) -> Dict[str, List[List[Any]]]:
    try:
        import openpyxl
    except ImportError:
        die("openpyxl required: pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    out: Dict[str, List[List[Any]]] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [list(r) for r in ws.iter_rows(values_only=True)]
    return out


def load_csv_rows(path: Path) -> List[List[Any]]:
    text = path.read_text(encoding="utf-8-sig")
    return list(csv.reader(text.splitlines()))


def load_json_annotated_csv(path: Path) -> List[List[Any]]:
    raw = json.loads(path.read_text(encoding="utf-8"))

    def find(o: Any) -> Optional[str]:
        if isinstance(o, dict):
            if isinstance(o.get("annotated_csv"), str):
                return o["annotated_csv"]
            for v in o.values():
                r = find(v)
                if r:
                    return r
        elif isinstance(o, list):
            for v in o:
                r = find(v)
                if r:
                    return r
        return None

    csv_text = find(raw)
    if not csv_text:
        die(f"no annotated_csv in {path}")
    return list(csv.reader(csv_text.splitlines()))


def discover_schema(
    rows: List[Sequence[Any]],
    base_col: str,
    new_col: str,
    sheet: Optional[str] = None,
) -> Schema:
    hi, headers = find_best_header(rows, base_col, new_col)
    ib = col_match(headers, base_col)
    inew = col_match(headers, new_col)
    if ib is None or inew is None:
        raise SchemaError(
            f"cannot locate product columns base={base_col!r} new={new_col!r}; "
            f"headers={[h for h in headers if h][:30]}"
        )
    id_cols = classify_id_columns(headers, ib, inew)
    # Prefer pad+pkg_pin; drop weak gpio labels when pad/pin already present
    roles = {r for r, _ in id_cols}
    if "pad" in roles and ("pin" in roles or "pkg_pin" in roles):
        id_cols = [(r, i) for r, i in id_cols if r in ("pad", "pin", "pkg_pin")]
    sch = Schema(hi, headers, ib, inew, id_cols, sheet=sheet)
    sch.notes = [
        f"header_row={hi}",
        f"kind_hint={sch.kind_hint()}",
        f"id_cols={id_cols}",
        f"base_col_idx={ib} new_col_idx={inew}",
    ]
    return sch


def _sheet_score(name: str, sch: Schema, base_col: str, new_col: str) -> int:
    score = 10 + 3 * len(sch.id_cols)
    nl = name.lower()
    # prefer sheets whose name contains the product tokens; do not boost
    # generic "gpio"/"test report" pages (those often false-match).
    bc = re.sub(r"[^a-z0-9]", "", base_col.lower())
    nc = re.sub(r"[^a-z0-9]", "", new_col.lower())
    nlc = re.sub(r"[^a-z0-9]", "", nl)
    if nc and nc in nlc:
        score += 8
    if bc and bc in nlc:
        score += 4
    if "配置" in name or "config" in nl:
        score += 3
    if any(x in nl for x in ("test report", "测试报告", "说明", "readme")):
        score -= 12
    bh = sch.headers[sch.base_col] if sch.base_col < len(sch.headers) else ""
    nh = sch.headers[sch.new_col] if sch.new_col < len(sch.headers) else ""
    if not bh or not nh:
        score -= 20
    return score


def pick_best_sheet(
    sheets: Dict[str, List[List[Any]]],
    base_col: str,
    new_col: str,
    prefer: Optional[str],
    explain: bool = False,
) -> Tuple[str, List[List[Any]], Schema]:
    if prefer:
        if prefer not in sheets:
            die(f"sheet not found: {prefer}; available={list(sheets)}")
        rows = sheets[prefer]
        try:
            return prefer, rows, discover_schema(rows, base_col, new_col, prefer)
        except SchemaError as e:
            die(str(e))

    ranked: List[Tuple[int, str, List[List[Any]], Schema]] = []
    errors: List[str] = []
    for name, rows in sheets.items():
        try:
            sch = discover_schema(rows, base_col, new_col, name)
            ranked.append((_sheet_score(name, sch, base_col, new_col), name, rows, sch))
        except SchemaError as e:
            errors.append(f"{name}: {e}")
    ranked.sort(key=lambda x: x[0], reverse=True)
    if explain:
        if ranked:
            print("  NOTE sheet candidates (score name):")
            for sc, name, _, _ in ranked:
                print(f"    {sc:4d}  {name}")
        if errors:
            print("  NOTE sheets skipped:")
            for e in errors[:12]:
                print(f"    {e}")
    if not ranked:
        die("no sheet contains BASE/NEW GPIO columns; tried:\n  " + "\n  ".join(errors[:12]))
    if len(ranked) >= 2 and ranked[0][0] - ranked[1][0] < 4:
        listing = "\n".join(f"  {sc}  {name}" for sc, name, _, _ in ranked[:8])
        die(
            "multiple GPIO sheets look valid; pass --sheet (or fill §2 工作表名).\n"
            + listing
        )
    best = ranked[0]
    return best[1], best[2], best[3]

def sig_core(s: str) -> str:
    """Strip annotations so PMU_IRQ vs PMIC_IRQ still differ, but EINT-only does not."""
    t = norm(s).lower()
    t = re.sub(r"\([^)]*\)", "", t)
    t = re.sub(r"\beint\d*\b", "", t)
    t = re.sub(r"[\s_\-/]+", "", t)
    return t


def classify_kind(base_sig: str, new_sig: str) -> str:
    ca, cb = sig_core(base_sig), sig_core(new_sig)
    if ca and cb and ca == cb:
        return "alias"
    return "functional"


def diffs_from_schema(rows: List[Sequence[Any]], sch: Schema) -> List[DiffRow]:
    out: List[DiffRow] = []
    need = max([sch.base_col, sch.new_col] + [i for _, i in sch.id_cols])
    for r in rows[sch.header_row + 1 :]:
        if not r or len(r) <= need:
            continue
        a, b = norm(r[sch.base_col]), norm(r[sch.new_col])
        if a == b:
            continue
        ids: Dict[str, str] = {}
        empty_ids = True
        for role, idx in sch.id_cols:
            val = norm(r[idx])
            if role == "pin" and val:
                val = val.split()[0]
            ids[role] = val
            if val:
                empty_ids = False
        if empty_ids:
            continue
        out.append(DiffRow(ids, a, b, classify_kind(a, b)))
    return out


def main() -> None:
    ap = argparse.ArgumentParser(
        description="Phase2 GPIO diff — dynamically discover table schema (examples A/B/C not required)"
    )
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--xlsx", type=Path, help="local xlsx path")
    src.add_argument("--csv", type=Path, help="local csv path")
    src.add_argument("--json-sheet", type=Path, help="sheets export json with annotated_csv")
    ap.add_argument("--sheet", default=None, help="xlsx sheet name; omit to auto-pick")
    ap.add_argument("--base-col", required=True, help="BASE product column name (fuzzy)")
    ap.add_argument("--new-col", required=True, help="NEW product column name (fuzzy)")
    ap.add_argument("--json-out", action="store_true", help="emit JSON")
    ap.add_argument("--explain", action="store_true", help="print discovered schema notes")
    args = ap.parse_args()

    sheet_name: Optional[str] = args.sheet
    try:
        if args.xlsx:
            sheets = load_xlsx_all_sheets(args.xlsx)
            sheet_name, rows, sch = pick_best_sheet(
                sheets, args.base_col, args.new_col, args.sheet, explain=args.explain
            )
        elif args.csv:
            rows = load_csv_rows(args.csv)
            sch = discover_schema(rows, args.base_col, args.new_col)
        else:
            rows = load_json_annotated_csv(args.json_sheet)
            sch = discover_schema(rows, args.base_col, args.new_col)
    except SchemaError as e:
        die(str(e))

    diffs = diffs_from_schema(rows, sch)
    n_alias = sum(1 for d in diffs if d.kind == "alias")
    payload = {
        "schema": {
            "sheet": sheet_name or sch.sheet,
            "kind_hint": sch.kind_hint(),
            "header_row": sch.header_row,
            "headers": [h for h in sch.headers if h][:40],
            "id_cols": [{"role": r, "index": i, "header": sch.headers[i]} for r, i in sch.id_cols],
            "base_col": {"index": sch.base_col, "header": sch.headers[sch.base_col]},
            "new_col": {"index": sch.new_col, "header": sch.headers[sch.new_col]},
            "notes": sch.notes,
        },
        "count": len(diffs),
        "alias_count": n_alias,
        "functional_count": len(diffs) - n_alias,
        "diffs": [d.as_dict() for d in diffs],
    }

    if args.json_out:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return

    print(
        f"sheet={payload['schema']['sheet']!r} kind_hint={sch.kind_hint()} "
        f"diffs={len(diffs)} alias={payload['alias_count']} "
        f"functional={payload['functional_count']} "
        f"id_cols={[r for r, _ in sch.id_cols]}"
    )
    if args.explain:
        for n in sch.notes:
            print(f"  NOTE {n}")
        print(f"  headers={payload['schema']['headers']}")
    for d in diffs:
        tag = "仅清单" if d.kind == "alias" else "功能"
        print(f"  [{tag}] {d.label()}: {d.base_sig!r} -> {d.new_sig!r}")


if __name__ == "__main__":
    main()
